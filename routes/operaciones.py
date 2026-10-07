from datetime import datetime
import math

import pytz
from flask import Blueprint, flash, redirect, render_template, request, url_for
from flask_login import current_user, login_required
from routes.decorators import permiso_required
from models import BalonGas, db

operaciones_bp = Blueprint('operaciones', __name__, url_prefix='/operaciones')
PERU_TZ = pytz.timezone('America/Lima')

def now_peru():
    return datetime.now(PERU_TZ).replace(tzinfo=None)

def hoy_peru():
    return datetime.now(PERU_TZ).date()

# ─── BALONES DE GAS ─────────────────────────────────────────
@operaciones_bp.route('/gas', methods=['GET', 'POST'])
@login_required
@permiso_required('gas')
def gas():
    if request.method == 'POST':
        accion = request.form.get('accion', 'nuevo')

        if accion == 'nuevo':
            def parse_date(s):
                try: return datetime.strptime(s, '%Y-%m-%d').date()
                except: return None
            try:
                cantidad = int(request.form.get('cantidad', 1) or 1)
                precio_unit = float(request.form.get('precio_unitario', 0) or 0)
                peso = float(request.form.get('peso_kg', 10) or 10)
                if cantidad < 1 or precio_unit < 0 or peso <= 0 or not all(map(math.isfinite, (precio_unit, peso))):
                    raise ValueError
                estado_ini = request.form.get('estado', 'disponible')
                if estado_ini not in {'disponible', 'en_uso'} or (estado_ini == 'en_uso' and cantidad != 1):
                    raise ValueError
                if estado_ini == 'en_uso' and BalonGas.query.filter_by(estado='en_uso').first():
                    flash('Ya hay un balón marcado como en uso. Agótalo antes de registrar otro en uso.', 'error')
                    return redirect(url_for('operaciones.gas'))
            except (TypeError, ValueError, OverflowError):
                flash('Cantidad, precio y peso deben ser valores válidos.', 'error')
                return redirect(url_for('operaciones.gas'))
            fecha_c       = parse_date(request.form.get('fecha_compra')) or hoy_peru()
            fecha_ini     = parse_date(request.form.get('fecha_inicio'))
            proveedor     = request.form.get('proveedor', '').strip()
            obs           = request.form.get('observaciones', '').strip()
            for _ in range(cantidad):
                b = BalonGas(
                    fecha_compra  = fecha_c,
                    fecha_inicio  = fecha_ini,
                    proveedor     = proveedor,
                    precio        = precio_unit,
                    peso_kg       = peso,
                    estado        = estado_ini,
                    observaciones = obs,
                    usuario_id    = current_user.id,
                    creado_en     = now_peru()
                )
                db.session.add(b)
            db.session.commit()
            total = round(cantidad * precio_unit, 2)
            flash(f'{cantidad} balón(es) registrado(s). Total: S/.{total:.2f}', 'success')

        elif accion == 'usar':
            b = BalonGas.query.get_or_404(int(request.form.get('balon_id')))
            if b.estado == 'agotado':
                flash('No se puede poner en uso un balón que ya fue agotado.', 'error')
                return redirect(url_for('operaciones.gas'))
            if b.estado != 'disponible':
                flash('Este balón ya está marcado como en uso.', 'error')
                return redirect(url_for('operaciones.gas'))
            otro_en_uso = BalonGas.query.filter(
                BalonGas.estado == 'en_uso', BalonGas.id != b.id
            ).first()
            if otro_en_uso:
                flash('Ya hay otro balón marcado como en uso. Agótalo antes de iniciar uno nuevo.', 'error')
                return redirect(url_for('operaciones.gas'))
            b.estado       = 'en_uso'
            b.fecha_inicio = hoy_peru()
            db.session.commit()
            flash('Balón marcado como en uso.', 'success')

        elif accion == 'agotar':
            b = BalonGas.query.get_or_404(int(request.form.get('balon_id')))
            b.estado   = 'agotado'
            b.fecha_fin = hoy_peru()
            if b.fecha_inicio:
                b.dias_uso = (b.fecha_fin - b.fecha_inicio).days or 1
            db.session.commit()
            flash(f'Balón cerrado. Duró {b.dias_uso or "?"} días.', 'success')

        elif accion == 'editar':
            def parse_date2(s):
                try: return datetime.strptime(s, '%Y-%m-%d').date()
                except: return None
            b = BalonGas.query.get_or_404(int(request.form.get('balon_id')))
            try:
                precio = float(request.form.get('precio_unitario', 0) or 0)
                peso = float(request.form.get('peso_kg', 10) or 10)
                if precio < 0 or peso <= 0 or not all(map(math.isfinite, (precio, peso))):
                    raise ValueError
            except (TypeError, ValueError, OverflowError):
                flash('Precio y peso deben ser valores válidos.', 'error')
                return redirect(url_for('operaciones.gas'))
            b.fecha_compra  = parse_date2(request.form.get('fecha_compra')) or b.fecha_compra
            b.proveedor     = request.form.get('proveedor', '').strip()
            b.precio        = precio
            b.peso_kg       = peso
            b.observaciones = request.form.get('observaciones', '').strip()
            db.session.commit()
            flash('Balón actualizado.', 'success')

        return redirect(url_for('operaciones.gas'))

    # Exportar a Excel (.xlsx)
    if request.args.get('export') == 'excel':
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from flask import make_response

        balones_exp = BalonGas.query.order_by(BalonGas.fecha_compra.desc()).all()
        wb = Workbook()
        ws = wb.active
        ws.title = 'Balones de Gas'

        # Estilos
        hdr_font  = Font(bold=True, color='FFFFFF', size=11)
        hdr_fill  = PatternFill('solid', fgColor='1E293B')
        hdr_align = Alignment(horizontal='center', vertical='center')
        thin = Side(style='thin', color='D1D5DB')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        headers = ['Fecha Compra', 'Peso (kg)', 'Precio Unit. S/.', 'Total S/.',
                   'Proveedor', 'Inicio Uso', 'Fin Uso', 'Días Uso', 'Estado', 'Observaciones']
        col_widths = [14, 10, 14, 12, 22, 12, 12, 10, 12, 30]

        # Título
        ws.merge_cells('A1:J1')
        ws['A1'] = 'REGISTRO DE BALONES DE GAS — Restaurante Turístico Marangani'
        ws['A1'].font = Font(bold=True, size=13, color='1E293B')
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.row_dimensions[1].height = 28

        # Cabecera
        for col, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=2, column=col, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = hdr_align
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = w
        ws.row_dimensions[2].height = 20

        # Datos
        estado_labels = {'disponible': 'Disponible', 'en_uso': 'En uso', 'agotado': 'Agotado'}
        fill_uso  = PatternFill('solid', fgColor='FEF3C7')
        fill_disp = PatternFill('solid', fgColor='DBEAFE')
        fill_agot = PatternFill('solid', fgColor='F1F5F9')

        for row_idx, b in enumerate(balones_exp, 3):
            total = round((b.precio or 0), 2)
            valores = [
                b.fecha_compra.strftime('%d/%m/%Y'),
                b.peso_kg,
                round(b.precio, 2) if b.precio else 0,
                total,
                b.proveedor or '',
                b.fecha_inicio.strftime('%d/%m/%Y') if b.fecha_inicio else '',
                b.fecha_fin.strftime('%d/%m/%Y') if b.fecha_fin else '',
                b.dias_uso or '',
                estado_labels.get(b.estado, b.estado),
                b.observaciones or ''
            ]
            fill_row = fill_uso if b.estado == 'en_uso' else (fill_disp if b.estado == 'disponible' else fill_agot)
            for col, val in enumerate(valores, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = border
                cell.alignment = Alignment(vertical='center',
                    horizontal='right' if col in (2,3,4,8) else 'center' if col in (6,7,9) else 'left')
                if col in (1,5,6,7,9,10):
                    cell.fill = fill_row
                ws.row_dimensions[row_idx].height = 16

        # Totales al pie
        n = len(balones_exp) + 3
        ws.cell(row=n, column=1, value='TOTAL').font = Font(bold=True)
        ws.cell(row=n, column=4,
            value=f'=SUM(D3:D{n-1})').font = Font(bold=True, color='166534')
        ws.cell(row=n, column=4).number_format = '"S/."#,##0.00'

        # Congelar cabecera
        ws.freeze_panes = 'A3'

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        resp = make_response(output.read())
        resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        resp.headers['Content-Disposition'] = 'attachment; filename=balones_gas.xlsx'
        return resp

    from sqlalchemy import func
    q_balones = BalonGas.query.order_by(BalonGas.fecha_compra.desc(), BalonGas.id.desc())
    pagina = q_balones.paginate(page=request.args.get('page', 1, type=int), per_page=50, error_out=False)
    balones = pagina.items
    prom_raw = db.session.query(func.avg(BalonGas.dias_uso)).filter(
        BalonGas.estado == 'agotado', BalonGas.dias_uso.isnot(None)
    ).scalar()
    prom_dias = round(float(prom_raw), 1) if prom_raw is not None else None
    en_uso_count = BalonGas.query.filter_by(estado='en_uso').count()
    en_uso = BalonGas.query.filter_by(estado='en_uso').order_by(
        BalonGas.fecha_inicio.desc(), BalonGas.id.desc()
    ).limit(1).all()
    disponibles_count = BalonGas.query.filter_by(estado='disponible').count()
    balones_total = q_balones.count()
    # Proyección: si hay uno en uso, cuántos días le quedan aprox
    dias_transcurridos = None
    dias_restantes_prom = None
    if en_uso and prom_dias and en_uso[0].fecha_inicio:
        dias_transcurridos  = (hoy_peru() - en_uso[0].fecha_inicio).days
        dias_restantes_prom = max(0, round(prom_dias - dias_transcurridos, 1))

    return render_template('operaciones/gas.html',
        balones=balones, prom_dias=prom_dias, en_uso=en_uso,
        pagina=pagina, en_uso_count=en_uso_count,
        disponibles_count=disponibles_count, balones_total=balones_total,
        dias_transcurridos=dias_transcurridos,
        dias_restantes_prom=dias_restantes_prom,
        hoy=hoy_peru())


